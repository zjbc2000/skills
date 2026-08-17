"""docx 解析器（v0.13）：段落 + 表格抽取。"""
from __future__ import annotations

import pytest

from backend.doc_loader.base import LoaderError, doc_to_text
from backend.doc_loader.docx_local import extract_text


@pytest.fixture
def docx_path(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("张三 - 高级工程师")
    d.add_paragraph("熟悉 RAG 与多轮对话")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "技能"
    table.cell(0, 1).text = "LangChain"
    p = tmp_path / "resume.docx"
    d.save(str(p))
    return p


def test_extract_text_paragraphs_and_tables(docx_path):
    text = extract_text(docx_path)
    assert "张三 - 高级工程师" in text
    assert "技能\tLangChain" in text


def test_doc_to_text_falls_through_to_docx(docx_path):
    doc = doc_to_text(docx_path, ["local_pymupdf", "local_docx"])
    assert doc.loader == "local_docx"
    assert "张三" in doc.text


def test_bad_docx_raises(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a docx")
    with pytest.raises(LoaderError):
        extract_text(bad)
