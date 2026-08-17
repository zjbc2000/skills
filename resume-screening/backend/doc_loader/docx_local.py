"""本地 docx 直读：python-docx 抽取段落与表格文本。"""
from __future__ import annotations

from .base import LoaderError


def extract_text(path) -> str:
    try:
        import docx
    except ImportError as e:
        raise LoaderError(f"python-docx 未安装: {e}") from e
    try:
        d = docx.Document(str(path))
    except Exception as e:
        raise LoaderError(f"无法打开 docx: {e}") from e

    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    text = "\n".join(parts).strip()
    if not text:
        raise LoaderError("docx 抽取文本为空")
    return text
